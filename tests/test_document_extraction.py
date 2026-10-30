"""Document extraction (WBS 7.4). Test files are generated in memory; OCR is faked
except in the opt-in real PaddleOCR test at the bottom."""

import io
import logging
import os
import sys
import time
import zipfile
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from PIL import Image, ImageDraw, ImageFont
from pydantic import ValidationError

from app.core.config import Settings
from app.document_processing import docx as docx_module
from app.document_processing import pdf as pdf_module
from app.document_processing import (
    CorruptDocumentError,
    DocumentError,
    DocumentTooLargeError,
    OcrUnavailableError,
    UnsupportedFileTypeError,
    detect_file_type,
    extract_document,
)
from tests.conftest import ServerThread, free_port
from app.document_processing.ocr import (
    DisabledOcrEngine,
    OcrResult,
    PaddleOcrEngine,
    SuryaOcrEngine,
    get_ocr_engine,
)

SECRET = "Student Dara scored 87 in Korean"


class FakeOcr:
    def __init__(self, text: str = "scanned words", confidence: float = 0.9) -> None:
        self.text = text
        self.confidence = confidence
        self.calls: list[Image.Image] = []

    def recognize(self, image: Image.Image) -> OcrResult:
        self.calls.append(image)
        return OcrResult(text=self.text, confidence=self.confidence)


def settings(**overrides) -> Settings:
    return Settings(_env_file=None, **overrides)


def make_pdf(pages: list[str | None]) -> bytes:
    """Minimal PDF; each page has the given text drawn in Helvetica, or nothing if None."""
    count = len(pages)
    font_id = 3 + 2 * count
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids ["
        + b" ".join(f"{3 + 2 * i} 0 R".encode() for i in range(count))
        + f"] /Count {count} >>".encode(),
    ]
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {4 + 2 * i} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_image(fmt: str, text: str | None = None) -> bytes:
    image = Image.new("RGB", (800, 200), "white")
    if text:
        ImageDraw.Draw(image).text((20, 60), text, fill="black", font=ImageFont.load_default(size=48))
    out = io.BytesIO()
    image.save(out, fmt)
    return out.getvalue()


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Attendance Policy", level=1)
    document.add_paragraph(SECRET)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Rate"
    table.cell(1, 0).text, table.cell(1, 1).text = "Dara", "95%"
    document.add_paragraph("Closing note")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --- type detection -------------------------------------------------------------


def test_detects_types_from_content_not_name():
    assert detect_file_type(make_pdf(["hello"])) == "pdf"
    assert detect_file_type(make_docx()) == "docx"
    assert detect_file_type(make_image("PNG")) == "png"
    assert detect_file_type(make_image("JPEG")) == "jpeg"
    assert detect_file_type(b"just some text") is None


def test_png_containing_pdf_marker_is_still_png():
    data = make_image("PNG")
    assert detect_file_type(data[:40] + b"%PDF-" + data[40:]) == "png"


def test_zip_that_is_not_docx_is_unsupported():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")  # looks like XLSX
    with pytest.raises(UnsupportedFileTypeError):
        extract_document(out.getvalue(), settings(), FakeOcr())


# --- PDF ------------------------------------------------------------------------


def test_text_pdf_uses_text_layer_with_page_numbers():
    ocr = FakeOcr()
    doc = extract_document(
        make_pdf(["First page about attendance rules", "Second page about allowance rates"]),
        settings(),
        ocr,
    )
    assert doc.file_type == "pdf"
    assert [(p.page_number, p.method) for p in doc.pages] == [(1, "text"), (2, "text")]
    assert doc.pages[0].text == "First page about attendance rules"
    assert doc.pages[1].text == "Second page about allowance rates"
    assert ocr.calls == []


def test_mixed_pdf_ocrs_only_pages_without_text():
    ocr = FakeOcr(text="text read by OCR", confidence=0.8)
    doc = extract_document(make_pdf(["A page with a real text layer", None]), settings(), ocr)
    assert [p.method for p in doc.pages] == ["text", "ocr"]
    assert doc.pages[1].text == "text read by OCR"
    assert doc.pages[1].ocr_confidence == 0.8
    assert len(ocr.calls) == 1
    # Rendered at the configured DPI: US Letter at 200 DPI is about 1700 x 2200 px.
    assert ocr.calls[0].size == pytest.approx((1700, 2200), abs=2)


def test_scanned_pdf_made_from_an_image_goes_to_ocr():
    scanned = io.BytesIO()
    Image.new("RGB", (850, 1100), "white").save(scanned, "PDF")
    ocr = FakeOcr()
    doc = extract_document(scanned.getvalue(), settings(), ocr)
    assert [p.method for p in doc.pages] == ["ocr"]
    assert len(ocr.calls) == 1


def test_blank_ocr_page_adds_warning():
    doc = extract_document(make_pdf([None]), settings(), FakeOcr(text="", confidence=None))
    assert doc.warnings == ["page 1: no text found"]


def test_corrupt_pdf_raises():
    with pytest.raises(CorruptDocumentError):
        extract_document(b"%PDF-1.4\nthis is not really a pdf", settings(), FakeOcr())


def test_pdf_over_page_limit_raises():
    with pytest.raises(DocumentTooLargeError):
        extract_document(make_pdf(["one", "two"]), settings(max_pdf_pages=1), FakeOcr())


def test_pdf_with_too_many_scanned_pages_raises():
    ocr = FakeOcr()
    with pytest.raises(DocumentTooLargeError):
        extract_document(make_pdf([None, None, None]), settings(max_ocr_pages=2), ocr)
    assert len(ocr.calls) == 2  # stops instead of OCR-ing the rest


def test_huge_pdf_page_render_is_capped(monkeypatch):
    monkeypatch.setattr(pdf_module, "MAX_RENDER_PIXELS", 1_000_000)
    ocr = FakeOcr()
    extract_document(make_pdf([None]), settings(), ocr)
    width, height = ocr.calls[0].size
    # pdfium rounds each side up to a whole pixel, so allow a fraction of a percent.
    assert 950_000 < width * height <= 1_005_000


# --- DOCX -----------------------------------------------------------------------


def test_docx_keeps_reading_order_headings_and_tables():
    doc = extract_document(make_docx(), settings(), FakeOcr())
    assert doc.file_type == "docx"
    assert len(doc.pages) == 1 and doc.pages[0].method == "text"
    assert doc.text == (
        "# Attendance Policy\n\n"
        f"{SECRET}\n\n"
        "Name | Rate\nDara | 95%\n\n"
        "Closing note"
    )


def test_docx_merged_cells_kept_once_but_equal_values_kept():
    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Week 1"
    table.cell(0, 2).text = "Week 2"
    for col in range(3):
        table.cell(1, col).text = "Yes"
    out = io.BytesIO()
    document.save(out)
    doc = extract_document(out.getvalue(), settings(), FakeOcr())
    assert doc.text == "Week 1 | Week 2\nYes | Yes | Yes"


def test_docx_with_invalid_xml_raises():
    # A valid DOCX whose main XML part is replaced, so parsing (not packaging) fails.
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(make_docx())) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = b"<not valid xml" if item.filename == "word/document.xml" else src.read(item)
            dst.writestr(item, body)
    with pytest.raises(CorruptDocumentError):
        extract_document(out.getvalue(), settings(), FakeOcr())


def test_docx_zip_bomb_rejected_before_parsing(monkeypatch):
    monkeypatch.setattr(docx_module, "MAX_UNZIPPED_BYTES", 1000)
    with pytest.raises(DocumentTooLargeError):
        extract_document(make_docx(), settings(), FakeOcr())


# --- images ---------------------------------------------------------------------


@pytest.mark.parametrize("fmt, file_type", [("PNG", "png"), ("JPEG", "jpeg")])
def test_images_go_to_ocr(fmt, file_type):
    ocr = FakeOcr(text="  photo text \n", confidence=0.75)
    doc = extract_document(make_image(fmt), settings(), ocr)
    assert doc.file_type == file_type
    assert doc.pages[0].text == "photo text"
    assert doc.pages[0].method == "ocr"
    assert doc.pages[0].ocr_confidence == 0.75


def test_truncated_image_raises():
    with pytest.raises(CorruptDocumentError):
        extract_document(make_image("PNG")[:60], settings(), FakeOcr())


def test_image_over_pixel_limit_raises():
    out = io.BytesIO()
    Image.new("1", (8000, 8000)).save(out, "PNG")  # 64 MP, but only a few KB on disk
    ocr = FakeOcr()
    with pytest.raises(DocumentTooLargeError):
        extract_document(out.getvalue(), settings(), ocr)
    assert ocr.calls == []


def test_ocr_disabled_raises_clear_error():
    with pytest.raises(OcrUnavailableError):
        extract_document(make_image("PNG"), settings(), DisabledOcrEngine())


def test_paddle_not_installed_raises_clear_error(monkeypatch):
    # Simulates a plain `poetry install` without the ocr group (no paddleocr, no numpy).
    monkeypatch.setitem(sys.modules, "paddleocr", None)
    monkeypatch.setitem(sys.modules, "numpy", None)
    engine = PaddleOcrEngine(language="en")
    for _ in range(2):  # the second call re-raises without retrying the import
        with pytest.raises(OcrUnavailableError, match="not installed"):
            extract_document(make_image("PNG"), settings(), engine)


# --- limits and logging ---------------------------------------------------------


def test_empty_and_oversized_files_raise():
    with pytest.raises(CorruptDocumentError):
        extract_document(b"", settings(), FakeOcr())
    with pytest.raises(DocumentTooLargeError):
        extract_document(b"%PDF-" + b"0" * (1024 * 1024), settings(max_document_mb=1), FakeOcr())


def test_surya_not_installed_raises_clear_error(monkeypatch):
    monkeypatch.setitem(sys.modules, "surya", None)
    engine = SuryaOcrEngine(inference_url=None, api_key="EMPTY")
    with pytest.raises(OcrUnavailableError, match="Surya is not installed"):
        extract_document(make_image("PNG"), settings(), engine)


def test_ocr_engine_is_chosen_by_setting():
    assert isinstance(get_ocr_engine(settings(ocr_engine="paddle")), PaddleOcrEngine)
    assert isinstance(get_ocr_engine(settings(ocr_engine="surya")), SuryaOcrEngine)
    assert isinstance(get_ocr_engine(settings(ocr_engine="none")), DisabledOcrEngine)


def test_surya_wrong_api_key_fails_clearly():
    pytest.importorskip("surya")
    fake = FastAPI()

    @fake.get("/v1/models")
    async def models():
        return JSONResponse({"error": "Unauthorized"}, status_code=401)

    server = ServerThread(fake, free_port())
    server.start()
    try:
        engine = SuryaOcrEngine(inference_url=f"{server.url}/v1", api_key="wrong")
        with pytest.raises(OcrUnavailableError, match="rejected the API key"):
            extract_document(make_image("PNG"), settings(), engine)
    finally:
        server.stop()


def test_surya_empty_result_on_non_blank_page_is_an_error():
    # Surya turns failed server requests into a page with no blocks.
    pytest.importorskip("surya")
    engine = SuryaOcrEngine(inference_url=None, api_key="EMPTY")
    engine._model = lambda images: [SimpleNamespace(blocks=[])]
    with pytest.raises(DocumentError, match="no result"):
        extract_document(make_image("PNG", "Some visible text"), settings(), engine)
    blank = extract_document(make_image("PNG"), settings(), engine)
    assert blank.pages[0].text == "" and blank.warnings == ["page 1: no text found"]


def test_surya_blocks_become_text_with_mean_confidence():
    pytest.importorskip("surya")

    def block(html, confidence, skipped=False, error=False):
        return SimpleNamespace(html=html, confidence=confidence, skipped=skipped, error=error)

    page = SimpleNamespace(
        blocks=[
            block("<h1>ការចូលរៀន</h1>", 0.9),  # Khmer: "attendance"
            block("<p>Line one<br>Line two</p>", 0.7),
            block("<table><tr><th>Name</th><th>Rate</th></tr><tr><td>Dara</td><td>95%</td></tr></table>", 0.8),
            block("", 1.0, skipped=True),  # e.g. a picture
            block("", 0.0, error=True),
        ]
    )
    engine = SuryaOcrEngine(inference_url=None, api_key="EMPTY")
    engine._model = lambda images: [page]  # stands in for Surya's RecognitionPredictor
    doc = extract_document(make_image("PNG"), settings(), engine)
    assert doc.pages[0].text == "ការចូលរៀន\n\nLine one\nLine two\n\nName | Rate\nDara | 95%"
    assert doc.pages[0].ocr_confidence == pytest.approx(0.8)


def test_engine_runtime_errors_become_document_error():
    engine = SuryaOcrEngine(inference_url=None, api_key="EMPTY")

    def broken(images):
        raise RuntimeError("server at http://internal-host:9999 refused")

    engine._model = broken
    with pytest.raises(DocumentError, match="^OCR failed$"):
        extract_document(make_image("PNG"), settings(), engine)


def test_invalid_limits_are_rejected():
    with pytest.raises(ValidationError):
        settings(pdf_ocr_dpi=0)


def test_logs_never_contain_document_text(caplog):
    with caplog.at_level(logging.DEBUG):
        extract_document(make_docx(), settings(), FakeOcr())
        extract_document(make_pdf([None]), settings(), FakeOcr(text=SECRET))
        with pytest.raises(DocumentTooLargeError) as too_many:
            extract_document(make_pdf([None, None]), settings(max_ocr_pages=1), FakeOcr(text=SECRET))
    assert "Extracted docx" in caplog.text and "Extracted pdf" in caplog.text
    assert "Dara" not in caplog.text
    assert "Dara" not in str(too_many.value)


# --- real PaddleOCR (opt-in: slow, downloads the model on first run) ------------


@pytest.mark.skipif(os.environ.get("RUN_OCR_TESTS") != "1", reason="set RUN_OCR_TESTS=1")
def test_real_paddleocr_reads_rendered_text():
    pytest.importorskip("paddleocr")
    engine = PaddleOcrEngine(language="en")
    doc = extract_document(make_image("PNG", "HRD Intelligence 2026"), settings(), engine)
    print("OCR text:", repr(doc.pages[0].text), "confidence:", doc.pages[0].ocr_confidence)
    assert "intelligence" in doc.pages[0].text.lower()
    assert doc.pages[0].ocr_confidence is not None and doc.pages[0].ocr_confidence > 0.5


# --- real Surya (opt-in: needs an inference server or llama-server; ~2.7 GB model) --

# Rendered with macOS CoreText, because Pillow without libraqm cannot shape Khmer.
KHMER_FIXTURE = Path(__file__).parent / "fixtures" / "khmer_sample.png"
KHMER_LINES = ["មជ្ឈមណ្ឌលកូរ៉េ សហ្វវែរ អេច អ ឌី", "សិស្សមានវត្តមាន ៩៥ ភាគរយ", "Attendance 2026"]


@pytest.mark.skipif(os.environ.get("RUN_SURYA_TESTS") != "1", reason="set RUN_SURYA_TESTS=1")
def test_real_surya_reads_khmer():
    pytest.importorskip("surya")
    # Like the service: SURYA_INFERENCE_URL / SURYA_INFERENCE_API_KEY from env vars or .env.
    # Without a URL, Surya starts a local llama-server instead.
    configured = Settings()
    engine = SuryaOcrEngine(
        inference_url=configured.surya_inference_url,
        api_key=configured.surya_inference_api_key.get_secret_value(),
    )
    start = time.monotonic()
    doc = extract_document(KHMER_FIXTURE.read_bytes(), settings(), engine)
    print("server:", "external (SURYA_INFERENCE_URL)" if configured.surya_inference_url else "local llama-server")
    print(f"took {time.monotonic() - start:.1f}s, confidence {doc.pages[0].ocr_confidence}")
    print("OCR text:", repr(doc.pages[0].text))
    assert [line for line in doc.pages[0].text.splitlines() if line] == KHMER_LINES
